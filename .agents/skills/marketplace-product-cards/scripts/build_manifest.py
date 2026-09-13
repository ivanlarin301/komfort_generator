#!/usr/bin/env python3
"""Build a read-only catalog/image inventory for product-card generation."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
INVALID_WINDOWS_CHARS = re.compile(r'[\\/:*?"<>|]')

HEADER_ALIASES = {
    "code": {"код ах", "код", "артикул", "sku", "product code"},
    "name": {"наименование", "название", "товар", "product name"},
    "photo_url": {"ссылка на фото", "фото", "photo url", "image url"},
    "product_url": {"ссылка на товар", "страница товара", "product url"},
}


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", " ", text).strip().casefold()


def safe_folder_name(value: Any) -> str:
    text = INVALID_WINDOWS_CHARS.sub("_", str(value or "")).strip()
    return text[:120]


def catalog_run_dir(catalog: Path) -> str:
    """Return the stable, human-readable output folder for one input workbook."""
    return safe_folder_name(catalog.stem) or "catalog"


def canonical_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header_map(row: tuple[Any, ...]) -> dict[str, int]:
    normalized = [normalize_text(value) for value in row]
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for index, label in enumerate(normalized):
            if label in aliases:
                result[field] = index
                break
    return result


def inspect_workbook(path: Path) -> tuple[int, str, dict[str, int], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    best: tuple[int, str, dict[str, int], int] | None = None
    try:
        for sheet in workbook.worksheets:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            mapping = header_map(first_row)
            score = (4 if "name" in mapping else 0) + (3 if "code" in mapping else 0)
            score += int("photo_url" in mapping) + int("product_url" in mapping)
            candidate = (score, sheet.title, mapping, sheet.max_row)
            if best is None or candidate[0] > best[0]:
                best = candidate
    finally:
        workbook.close()
    if best is None:
        raise ValueError(f"Книга не содержит листов: {path}")
    return best


def select_catalog(project_root: Path, explicit: str | None) -> tuple[Path, str, dict[str, int]]:
    if explicit:
        catalog = Path(explicit)
        if not catalog.is_absolute():
            catalog = project_root / catalog
        candidates = [catalog.resolve()]
    else:
        candidates = [
            path.resolve()
            for path in project_root.glob("*.xlsx")
            if not path.name.startswith("~$")
        ]
    if not candidates:
        raise FileNotFoundError("Не найден подходящий .xlsx в корне проекта")

    ranked: list[tuple[int, int, Path, str, dict[str, int]]] = []
    errors: list[str] = []
    for path in candidates:
        if not path.is_file():
            errors.append(f"Файл не найден: {path}")
            continue
        try:
            score, sheet, mapping, rows = inspect_workbook(path)
            ranked.append((score, rows, path, sheet, mapping))
        except Exception as exc:  # keep evaluating other workbooks
            errors.append(f"{path.name}: {exc}")

    ranked.sort(key=lambda item: (item[0], item[1], item[2].stat().st_mtime), reverse=True)
    if not ranked or ranked[0][0] < 7:
        detail = "; ".join(errors) if errors else "не распознаны столбцы кода и наименования"
        raise ValueError(f"Не удалось выбрать каталог: {detail}")
    _, _, path, sheet, mapping = ranked[0]
    return path, sheet, mapping


def find_folder(name: str, folders: list[Path]) -> tuple[Path | None, str]:
    expected = normalize_text(safe_folder_name(name))
    exact = [folder for folder in folders if normalize_text(folder.name) == expected]
    if len(exact) == 1:
        return exact[0], "exact"

    scores = sorted(
        ((SequenceMatcher(None, expected, normalize_text(folder.name)).ratio(), folder) for folder in folders),
        key=lambda item: item[0],
        reverse=True,
    )
    if not scores or scores[0][0] < 0.92:
        return None, "missing"
    if len(scores) > 1 and scores[0][0] - scores[1][0] < 0.02:
        return None, "ambiguous"
    return scores[0][1], "fuzzy"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def inspect_image(path: Path, project_root: Path) -> dict[str, Any]:
    record: dict[str, Any] = {
        "path": path.relative_to(project_root).as_posix(),
        "bytes": path.stat().st_size,
    }
    try:
        with Image.open(path) as image:
            record.update(
                {
                    "format": image.format,
                    "width": image.width,
                    "height": image.height,
                    "frames": getattr(image, "n_frames", 1),
                    "sha256": file_sha256(path),
                    "readable": True,
                }
            )
    except Exception as exc:
        record.update({"readable": False, "error": str(exc)})
    return record


def build_manifest(project_root: Path, catalog_arg: str | None) -> dict[str, Any]:
    catalog, sheet_name, mapping = select_catalog(project_root, catalog_arg)
    images_root = project_root / "images"
    folders = sorted([path for path in images_root.iterdir() if path.is_dir()]) if images_root.is_dir() else []

    workbook = load_workbook(catalog, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    products: list[dict[str, Any]] = []
    try:
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            name = str(row[mapping["name"]] or "").strip()
            code = canonical_code(row[mapping["code"]])
            if not name and not code:
                continue
            folder, match_type = find_folder(name, folders)
            sources: list[dict[str, Any]] = []
            if folder:
                source_paths = sorted(
                    path for path in folder.iterdir() if path.is_file() and path.suffix.casefold() in IMAGE_EXTENSIONS
                )
                sources = [inspect_image(path, project_root) for path in source_paths]

            readable_count = sum(bool(source.get("readable")) for source in sources)
            if folder is None:
                source_status = match_type
            elif not sources:
                source_status = "empty"
            elif readable_count != len(sources):
                source_status = "damaged"
            else:
                source_status = "ready"

            product = {
                "row": row_number,
                "code": code,
                "name": name,
                "photo_url": str(row[mapping["photo_url"]] or "").strip() if "photo_url" in mapping else "",
                "product_url": str(row[mapping["product_url"]] or "").strip() if "product_url" in mapping else "",
                "source_folder": folder.relative_to(project_root).as_posix() if folder else None,
                "folder_match": match_type,
                "source_status": source_status,
                "sources": sources,
            }
            products.append(product)
    finally:
        workbook.close()

    duplicate_map: dict[str, list[str]] = {}
    for product in products:
        for source in product["sources"]:
            if source.get("readable") and source.get("sha256"):
                duplicate_map.setdefault(source["sha256"], []).append(source["path"])
    duplicates = [paths for paths in duplicate_map.values() if len(paths) > 1]

    return {
        "schema_version": 1,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "project_root": project_root.as_posix(),
        "catalog": {
            "path": catalog.relative_to(project_root).as_posix(),
            "sheet": sheet_name,
        },
        "run_dir": f"output/product-cards/{catalog_run_dir(catalog)}",
        "inventory": {
            "product_count": len(products),
            "ready_count": sum(product["source_status"] == "ready" for product in products),
            "image_count": sum(len(product["sources"]) for product in products),
            "duplicate_file_groups": duplicates,
        },
        "products": products,
        "cards": [],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".", help="Project directory; defaults to current directory")
    parser.add_argument("--catalog", help="Explicit .xlsx path, absolute or relative to project root")
    parser.add_argument("--output", help="Manifest path; by default it is isolated by input workbook")
    parser.add_argument("--dry-run", action="store_true", help="Print summary without writing a manifest")
    args = parser.parse_args()

    project_root = Path(args.project_root).resolve()
    try:
        manifest = build_manifest(project_root, args.catalog)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    output = Path(args.output or f"{manifest['run_dir']}/manifest.json")
    if not output.is_absolute():
        output = project_root / output

    if output.exists():
        try:
            previous = json.loads(output.read_text(encoding="utf-8"))
            if isinstance(previous.get("cards"), list):
                manifest["cards"] = previous["cards"]
        except Exception:
            pass

    summary = manifest["inventory"]
    print(f"Catalog: {manifest['catalog']['path']} [{manifest['catalog']['sheet']}]")
    print(
        f"Products: {summary['product_count']}; ready: {summary['ready_count']}; "
        f"images: {summary['image_count']}; duplicate groups: {len(summary['duplicate_file_groups'])}"
    )
    for product in manifest["products"]:
        if product["source_status"] != "ready":
            print(f"WARN {product['code']} {product['name']}: {product['source_status']}")

    if args.dry_run:
        return 0
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Manifest: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
