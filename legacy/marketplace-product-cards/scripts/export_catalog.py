#!/usr/bin/env python3
"""Create a catalog copy with traceable links to sources and generated cards."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font


INVALID_WINDOWS_CHARS = re.compile(r'[\\/:*?"<>|]')


def safe_name(value: str) -> str:
    return INVALID_WINDOWS_CHARS.sub("_", value).strip()[:120]


def link_cell(cell: Any, target: Path, label: str, workbook_dir: Path) -> None:
    relative = Path(__import__("os").path.relpath(target, workbook_dir)).as_posix()
    cell.value = label
    cell.hyperlink = relative
    cell.style = "Hyperlink"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, help="Path to the manifest for one catalog run")
    args = parser.parse_args()

    manifest_path = Path(args.manifest).resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    project_root = Path(manifest["project_root"])
    catalog_path = project_root / manifest["catalog"]["path"]
    run_dir = manifest_path.parent
    output_path = run_dir / f"{safe_name(catalog_path.stem)}_с_карточками.xlsx"
    shutil.copy2(catalog_path, output_path)

    card_by_code: dict[str, dict[str, Any]] = {}
    for card in manifest.get("cards", []):
        for code in card.get("product_codes", []):
            card_by_code[str(code)] = card

    products = {str(item.get("code", "")): item for item in manifest.get("products", [])}
    workbook = load_workbook(output_path)
    sheet = workbook[manifest["catalog"]["sheet"]]
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
    code_column = headers.get("Код АХ") or headers.get("Код") or headers.get("Артикул")
    if not code_column:
        raise ValueError("В каталоге не найдена колонка кода товара")

    extra_headers = [
        "Статус карточки",
        "Папка исходных фото",
        "Список исходных фото",
        "Первая исходная фотография",
        "Папка карточки",
        "Готовая карточка",
    ]
    start_column = sheet.max_column + 1
    for offset, label in enumerate(extra_headers):
        cell = sheet.cell(1, start_column + offset, label)
        cell.font = Font(bold=True)

    for row in range(2, sheet.max_row + 1):
        code = str(sheet.cell(row, code_column).value or "").strip()
        product = products.get(code)
        if not product:
            continue
        card = card_by_code.get(code, {})
        status = card.get("status") or ("blocked" if product.get("source_status") != "ready" else "planned")
        sheet.cell(row, start_column, status)
        source_folder = product.get("source_folder")
        if source_folder:
            source_dir = project_root / source_folder
            link_cell(sheet.cell(row, start_column + 1), source_dir, source_folder, output_path.parent)
            sources = product.get("sources", [])
            if sources:
                first_source = project_root / sources[0]["path"]
                source_list = "\n".join(source["path"] for source in sources)
                list_cell = sheet.cell(row, start_column + 2, source_list)
                list_cell.alignment = Alignment(wrap_text=True, vertical="top")
                link_cell(sheet.cell(row, start_column + 3), first_source, "Открыть первое фото", output_path.parent)
        result_path = card.get("result_path") or card.get("output_path")
        if result_path:
            card_path = project_root / result_path
            link_cell(sheet.cell(row, start_column + 4), card_path.parent, card_path.parent.relative_to(project_root).as_posix(), output_path.parent)
            link_cell(sheet.cell(row, start_column + 5), card_path, "Открыть карточку", output_path.parent)

    for column in range(start_column, start_column + len(extra_headers)):
        sheet.column_dimensions[__import__("openpyxl").utils.get_column_letter(column)].width = 24
    workbook.save(output_path)
    print(f"Output catalog: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
