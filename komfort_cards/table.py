"""Read the operator's Excel table exactly (no OCR): codes, names, prices and other columns (FR-2)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from komfort_cards.console import CardsError, say
from komfort_cards.formatting import normalize_code

ALIASES = {
    "code": ("код ах", "кодах", "код", "код товара", "артикул", "sku"),
    "name": ("наименование", "наименование продукта", "наименование товара", "название", "товар"),
}
PRICE_HINT = re.compile(r"цен|стоим|price", re.IGNORECASE)


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def read_table(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    if not path.is_file():
        raise CardsError(f"Файл таблицы не найден: {path}")
    if path.name.startswith("~$"):
        raise CardsError("Это временный файл Excel (~$...). Укажите сам файл таблицы.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        best = None
        for sheet in workbook.worksheets:
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(v not in (None, "") for v in r)]
            if not rows:
                continue
            headers = [str(v).strip() if v is not None else "" for v in rows[0]]
            score = sum(any(_norm(h) == a for h in headers for a in aliases) for aliases in ALIASES.values())
            if best is None or score > best[0]:
                best = (score, sheet.title, headers, rows[1:])
    finally:
        workbook.close()
    if best is None:
        raise CardsError("Таблица пустая.")
    _, sheet_name, headers, body = best
    columns = {key: next((i for i, h in enumerate(headers) if _norm(h) in aliases), None) for key, aliases in ALIASES.items()}
    price_columns = [h for h in headers if PRICE_HINT.search(h)]
    records = []
    for n, row in enumerate(body, start=1):
        record = {headers[i] or f"col{i + 1}": row[i] for i in range(min(len(headers), len(row)))}
        if columns["code"] is not None:
            record[headers[columns["code"]]] = normalize_code(row[columns["code"]])
        records.append({"n": n, **{k: ("" if v is None else v) for k, v in record.items()}})
    return {"file": str(path), "sheet": sheet_name, "headers": headers, "code_column": _header(headers, columns["code"]),
            "name_column": _header(headers, columns["name"]), "price_columns": price_columns, "rows": records}


def _header(headers: list[str], index: int | None) -> str | None:
    return headers[index] if index is not None else None


def cmd_table(args) -> dict[str, Any]:
    data = read_table(Path(args.path))
    headers = ["№"] + data["headers"]
    say(f"Лист «{data['sheet']}», строк: {len(data['rows'])}")
    say("| " + " | ".join(headers) + " |")
    say("|" + "---|" * len(headers))
    for row in data["rows"]:
        cells = [str(row["n"])] + [str(row.get(h or "", "")) for h in data["headers"]]
        say("| " + " | ".join(c.replace("|", "/") for c in cells) + " |")
    if len(data["price_columns"]) > 1:
        say(f"ВНИМАНИЕ: несколько колонок с ценой ({', '.join(data['price_columns'])}) — уточните у оператора, какую ставить.")
    if not data["code_column"] or not data["price_columns"]:
        say("ВНИМАНИЕ: не нашёл колонку кода или цены — проверьте заголовки таблицы.")
    return data
